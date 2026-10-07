"""AI Resume ATS Checker - Streamlit + Google Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.5-flash"
MAX_CHARS = 20000  # safety limit on resume / job description length

st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")


# ---------------------------------------------------------------- helpers
def get_secret_key() -> str:
    """Look for the API key in Streamlit secrets, then environment variables."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:  # no secrets.toml present
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def extract_text(uploaded_file) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")

    raise ValueError("Unsupported file type. Please upload PDF, DOCX or TXT.")


PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and senior recruiter.
Evaluate the resume below{jd_clause}.

Scoring rubric (each 0-100):
- keywords: relevant skills / keywords{jd_kw}
- formatting: ATS-friendly structure, standard section headings, readable text
- experience: impact, quantified achievements, action verbs
- skills: clarity and relevance of the skills section
- education: completeness and clarity of education / certifications
- overall: weighted overall ATS score

Be honest and strict; do not inflate scores. Base everything ONLY on the resume text.
Give 5-8 concrete, specific improvements (reference actual parts of the resume).

Return ONLY valid JSON with exactly this shape:
{{
  "overall_score": <int 0-100>,
  "summary": "<2-3 sentence assessment>",
  "section_scores": {{"keywords": <int>, "formatting": <int>, "experience": <int>, "skills": <int>, "education": <int>}},
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": [{{"priority": "High|Medium|Low", "issue": "...", "suggestion": "..."}}],
  "rewritten_summary": "<an improved professional summary for this candidate>"
}}

{jd_block}RESUME:
\"\"\"
{resume}
\"\"\"
"""


def build_prompt(resume: str, jd: str) -> str:
    if jd.strip():
        return PROMPT.format(
            jd_clause=" against the job description",
            jd_kw=" and match with the job description",
            jd_block=f'JOB DESCRIPTION:\n"""\n{jd[:MAX_CHARS]}\n"""\n\n',
            resume=resume[:MAX_CHARS],
        )
    return PROMPT.format(
        jd_clause=" for general ATS compatibility",
        jd_kw="",
        jd_block="",
        resume=resume[:MAX_CHARS],
    )


def parse_json(text: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences."""
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def analyze_resume(api_key: str, model: str, resume: str, jd: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume, jd),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    result = parse_json(response.text)
    if not isinstance(result, dict):
        raise ValueError("The model returned an unexpected response format.")
    return result


def score_color(score: int) -> str:
    return "🟢" if score >= 75 else "🟠" if score >= 50 else "🔴"


# --------------------------------------------------------------------- UI
st.title("📄 AI Resume ATS Checker")
st.caption("Upload your resume to get an ATS score and tips to improve it.")

with st.sidebar:
    st.header("⚙️ Settings")
    api_key = get_secret_key()
    if not api_key:
        api_key = st.text_input(
            "Gemini API key", type="password",
            help="Free key: https://aistudio.google.com/apikey",
        )
    else:
        st.success("API key loaded from secrets")
    model_name = st.text_input("Model", value=DEFAULT_MODEL,
                               help="e.g. gemini-3.5-flash or gemini-2.5-flash")
    st.markdown("---")
    st.caption("Your resume is sent to Google's Gemini API for analysis. "
               "Nothing is stored by this app.")

left, right = st.columns(2)
with left:
    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
with right:
    job_desc = st.text_area(
        "Job description (optional, for a targeted score)", height=180,
        placeholder="Paste the job description here...",
    )

if st.button("Analyze resume", type="primary", disabled=uploaded is None):
    if not api_key:
        st.error("Please provide a Gemini API key in the sidebar.")
        st.stop()
    try:
        resume_text = extract_text(uploaded).strip()
    except Exception as e:
        st.error(f"Could not read the file: {e}")
        st.stop()

    if len(resume_text) < 50:
        st.error("Couldn't extract enough text. If your PDF is a scanned image, "
                 "ATS systems can't read it either - export a text-based PDF or DOCX.")
        st.stop()

    with st.spinner("Analyzing your resume..."):
        try:
            result = analyze_resume(api_key, model_name.strip() or DEFAULT_MODEL,
                                    resume_text, job_desc)
        except json.JSONDecodeError:
            st.error("The AI returned an invalid response. Please try again.")
            st.stop()
        except Exception as e:
            st.error(f"Analysis failed: {e}")
            st.stop()

    overall = clamp(result.get("overall_score"))
    st.divider()
    c1, c2 = st.columns([1, 3])
    c1.metric("ATS Score", f"{overall}/100")
    c1.progress(overall / 100)
    c2.markdown(f"### {score_color(overall)} Summary")
    c2.write(result.get("summary", ""))

    scores = result.get("section_scores", {}) or {}
    if scores:
        st.subheader("Score breakdown")
        cols = st.columns(len(scores))
        for col, (label, val) in zip(cols, scores.items()):
            v = clamp(val)
            col.metric(label.capitalize(), f"{v}/100")
            col.progress(v / 100)

    a, b = st.columns(2)
    with a:
        st.subheader("✅ Strengths")
        for s in result.get("strengths", []) or []:
            st.markdown(f"- {s}")
    with b:
        st.subheader("⚠️ Weaknesses")
        for w in result.get("weaknesses", []) or []:
            st.markdown(f"- {w}")

    missing = result.get("missing_keywords", []) or []
    if missing:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in missing))

    st.subheader("🛠️ Suggested improvements")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for item in result.get("improvements", []) or []:
        if isinstance(item, dict):
            icon = icons.get(str(item.get("priority", "")).lower(), "⚪")
            with st.expander(f"{icon} {item.get('priority', '')}: {item.get('issue', '')}"):
                st.write(item.get("suggestion", ""))
        else:
            st.markdown(f"- {item}")

    if result.get("rewritten_summary"):
        st.subheader("✍️ Suggested professional summary")
        st.info(result["rewritten_summary"])

    st.download_button("Download report (JSON)", json.dumps(result, indent=2),
                       file_name="ats_report.json", mime="application/json")
